import io

import pytest
from docx import Document
from docx.shared import Cm

from app.utils.modul_docx import generate_modul_docx
from tests.unit.test_modul_pdf import make_content

DOCX_TABLE_COUNT = 7  # identity, Pancasila, one per meeting (2), overview, rubric, glossary


def open_docx(content=None):
    return Document(io.BytesIO(generate_modul_docx(content or make_content())))


def paragraphs(document, style_prefix=""):
    return [p.text for p in document.paragraphs if p.style.name.startswith(style_prefix)]


def table_rows(table):
    return [[cell.text for cell in row.cells] for row in table.rows]


def all_text(document):
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*paragraphs(document), *cells])


def test_the_file_is_a_valid_docx_package():
    data = generate_modul_docx(make_content())
    assert data.startswith(b"PK")  # a zip container
    assert open_docx().paragraphs


def test_identity_title_and_every_component_are_present():
    text = all_text(open_docx())
    for expected in (
        "MODUL AJAR KURIKULUM MERDEKA",
        "Pelaku dan Kegiatan Ekonomi di Sekitar Kita",
        "Kelas 9 (Fase D)",
        "2 pertemuan x 80 menit (160 menit)",
        "A. INFORMASI UMUM",
        "B. KOMPONEN INTI",
        "C. LAMPIRAN",
        "Tujuan pembelajaran",
        "Pertemuan 1",
        "Pertemuan 2",
        "Asesmen diagnostik",
        "Pengayaan",
        "Remedial",
        "LKPD: Peta alur barang",
        "Bahan bacaan peserta didik",
    ):
        assert expected in text, expected


def test_parts_and_sections_use_word_headings_so_the_document_has_an_outline():
    document = open_docx()
    assert paragraphs(document, "Heading 1") == ["A. INFORMASI UMUM", "B. KOMPONEN INTI", "C. LAMPIRAN"]
    sections = paragraphs(document, "Heading 2")
    assert "Tujuan pembelajaran" in sections and "Kegiatan pembelajaran" in sections and "Asesmen" in sections


def test_activities_assessment_and_glossary_are_real_tables():
    document = open_docx()
    assert len(document.tables) == DOCX_TABLE_COUNT

    by_header = {tuple(table.rows[0].cells[i].text for i in range(len(table.columns))): table for table in document.tables}
    activities = [t for t in document.tables if table_rows(t)[0] == ["Tahap", "Waktu (menit)", "Kegiatan", "Metode"]]
    assert len(activities) == 2  # one per meeting
    assert table_rows(activities[0])[1][:2] == ["Pendahuluan", "15"]
    assert ("Dimensi", "Wujud dalam kegiatan") in by_header
    assert ("Istilah", "Pengertian") in by_header
    assert any(header[0] == "Kriteria" for header in by_header)  # rubric


def test_the_header_row_of_a_data_table_repeats_across_pages_but_the_identity_table_has_none():
    document = open_docx()
    identity, *data_tables = document.tables

    def repeats(table):
        return table.rows[0]._tr.trPr is not None and table.rows[0]._tr.trPr.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
        ) is not None

    assert not repeats(identity)
    assert all(repeats(table) for table in data_tables)


def test_the_identity_table_pairs_each_label_with_its_value():
    identity = open_docx().tables[0]
    rows = dict(table_rows(identity))
    assert rows["Mata pelajaran"] == "IPS"
    assert rows["Topik"] == "Pelaku dan kegiatan ekonomi"
    assert rows["Model pembelajaran"] == "Diskusi kelompok"


def test_numbered_lists_start_again_at_one_and_cited_objectives_show_their_source():
    document = open_docx()
    numbered = [p for p in paragraphs(document) if p[:2] in ("1.", "2.", "3.")]

    assert numbered.count([n for n in numbered if n.startswith("1.")][0]) == 1
    assert sum(1 for p in numbered if p.startswith("1.")) == 2  # objectives and worksheet tasks both begin at 1
    assert any("menganalisis interaksi antarpelaku ekonomi. [S1]" in p for p in numbered)
    assert any(p.startswith("2.") and "[S1]" not in p for p in numbered)  # an ungrounded objective has no marker


def test_references_are_listed_with_their_pages():
    assert "[S1] CP IPS.pdf, hlm. 14-15, Elemen Pemahaman" in paragraphs(open_docx())


def test_typographic_punctuation_survives():
    content = make_content()
    content.title = "Ekonomi – “Pelaku” dan ‘Kegiatan’ … ≥ 80%"
    content.meetings[0].activities.rows[0][2] = "Guru bertanya: “Apa yang kamu beli hari ini?” – lalu mencatat."

    text = all_text(open_docx(content))

    assert "Ekonomi – “Pelaku” dan ‘Kegiatan’ … ≥ 80%" in text
    assert "“Apa yang kamu beli hari ini?” – lalu mencatat." in text


def test_control_characters_in_model_text_do_not_break_the_export():
    content = make_content()
    content.title = "Judul\x00 dengan\x0b karakter\x1f kontrol"
    content.meetings[0].activities.rows[0][2] = "Sel dengan \x08 backspace"
    content.attachments.reading_material = "Paragraf satu.\n\x0c\nParagraf dua."

    text = all_text(open_docx(content))

    assert "Judul dengan karakter kontrol" in text
    assert "Sel dengan  backspace" in text


def test_a_cell_with_several_lines_keeps_them():
    content = make_content()
    content.meetings[0].activities.rows[0][2] = "Baris satu\nBaris dua"

    activities = next(t for t in open_docx(content).tables if table_rows(t)[0][0] == "Tahap")

    assert "Baris satu" in table_rows(activities)[1][2] and "Baris dua" in table_rows(activities)[1][2]


def test_the_page_is_a_4_with_page_numbers_in_the_footer():
    document = open_docx()
    section = document.sections[0]

    assert abs(section.page_width - Cm(21.0)) < Cm(0.01) and abs(section.page_height - Cm(29.7)) < Cm(0.01)
    footer_xml = section.footer._element.xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml


def test_a_long_module_exports_all_of_its_meetings():
    document = open_docx(make_content(meetings=6))
    activities = [t for t in document.tables if table_rows(t)[0][0] == "Tahap"]
    assert len(activities) == 6


def test_optional_sections_may_be_empty():
    content = make_content()
    content.initial_competencies = []
    content.pancasila_profile = []
    content.facilities = []
    content.enrichment = []
    content.remedial = []
    content.reflection.teacher = []
    content.reflection.students = []
    content.attachments.glossary = []
    content.attachments.reading_material = ""
    content.references = []

    text = all_text(open_docx(content))

    assert "Pertemuan 1" in text
    assert "Pengayaan" not in text and "Glosarium" not in text and "Sumber rujukan" not in text


@pytest.mark.parametrize("title", ["", "   "])
def test_a_blank_title_falls_back_to_the_topic(title):
    content = make_content()
    content.title = title

    assert "Pelaku dan kegiatan ekonomi" in paragraphs(open_docx(content))

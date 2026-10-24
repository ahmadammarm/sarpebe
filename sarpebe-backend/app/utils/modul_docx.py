"""
Word export for a Modul Ajar in schema version 2. It follows the same order as the PDF export
(modul_pdf.py): identity, A. Informasi umum, B. Komponen inti, C. Lampiran and the reference list.
Headings use Word's heading styles so the document has a navigable outline, and tables repeat their
header row across pages.
"""
import re
from io import BytesIO

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table as DocxTable, _Cell
from docx.text.paragraph import Paragraph

from app.schemas.lesson_plan_content import ModulContent, Table
from app.utils.modul_pdf import ACTIVITY_COLUMN_WIDTHS, reference_line

FONT_NAME = "Calibri"
INK = RGBColor(0x1F, 0x2A, 0x44)
HEADER_FILL = "E8EEFC"
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
MARGIN_CM = 2.0
CONTENT_WIDTH_CM = PAGE_WIDTH_CM - 2 * MARGIN_CM
LIST_INDENT_CM = 0.7

# Characters Word's XML cannot hold (everything below space except tab, newline and carriage return).
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(text: str) -> str:
    """Model text can carry control characters or lone surrogates, which make python-docx raise."""
    return _INVALID_XML.sub("", text).encode("utf-8", "ignore").decode("utf-8")


# --- Document setup ---------------------------------------------------------------------------

def _style_fonts(document: DocumentType) -> None:
    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    for name, size, before, after in (("Heading 1", 14, 14, 4), ("Heading 2", 11.5, 10, 3)):
        style = document.styles[name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet",):
        document.styles[name].font.name = FONT_NAME


def _setup_page(document: DocumentType) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(MARGIN_CM))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Halaman ").font.size = Pt(8.5)
    _add_field(footer, "PAGE")
    footer.add_run(" dari ").font.size = Pt(8.5)
    _add_field(footer, "NUMPAGES")


def _add_field(paragraph: Paragraph, instruction: str) -> None:
    """A Word field (page number, page count) that Word fills in when it lays the document out."""
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    properties.append(size)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


# --- Text -------------------------------------------------------------------------------------

def _write_part(document: DocumentType, text: str) -> None:
    """A major part, e.g. 'B. KOMPONEN INTI'."""
    document.add_heading(_clean(text), level=1)


def _write_heading(document: DocumentType, text: str) -> None:
    document.add_heading(_clean(text), level=2)


def _write_label(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(_clean(text)).bold = True


def _write_paragraph(document: DocumentType, text: str) -> None:
    if text.strip():
        document.add_paragraph(_clean(text))


def _write_list(
    document: DocumentType,
    items: list[str],
    *,
    numbered: bool = False,
    suffixes: list[str] | None = None,
) -> None:
    """
    Bullets use Word's list style. Numbers are written out, because Word's numbered-list style
    carries one running count through the whole document instead of starting again per list.
    """
    for position, item in enumerate(items, start=1):
        if not item.strip():
            continue
        suffix = f" {suffixes[position - 1]}" if suffixes and suffixes[position - 1] else ""
        text = _clean(item + suffix)
        if numbered:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(LIST_INDENT_CM)
            paragraph.paragraph_format.first_line_indent = Cm(-LIST_INDENT_CM)
            paragraph.add_run(f"{position}.\t{text}")
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(LIST_INDENT_CM))
        else:
            document.add_paragraph(text, style="List Bullet")


# --- Tables -----------------------------------------------------------------------------------

def _shade(cell: _Cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _row_flag(row, tag: str) -> None:
    flag = OxmlElement(tag)
    flag.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(flag)


def _fill_cell(cell: _Cell, text: str, *, bold: bool = False, size: float = 9.5) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(_clean(text))
    run.font.size = Pt(size)
    run.bold = bold


def _apply_widths(table: DocxTable, fractions: tuple[float, ...] | None, columns: int) -> None:
    fractions = fractions if fractions and len(fractions) == columns else tuple(1 / columns for _ in range(columns))
    table.autofit = False
    for row in table.rows:
        for cell, fraction in zip(row.cells, fractions):
            cell.width = Cm(CONTENT_WIDTH_CM * fraction)


def _write_table(
    document: DocumentType,
    columns: list[str],
    rows: list[list[str]],
    *,
    caption: str = "",
    col_widths: tuple[float, ...] | None = None,
) -> None:
    if caption.strip():
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(_clean(caption))
        run.italic = True
        run.font.size = Pt(9.5)

    table = document.add_table(rows=1, cols=len(columns), style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header = table.rows[0]
    _row_flag(header, "w:tblHeader")  # repeat on every page
    for cell, column in zip(header.cells, columns):
        _fill_cell(cell, column, bold=True)
        _shade(cell, HEADER_FILL)

    for values in rows:
        row = table.add_row()
        _row_flag(row, "w:cantSplit")
        for cell, value in zip(row.cells, values):
            _fill_cell(cell, value)

    _apply_widths(table, col_widths, len(columns))
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _write_model_table(document: DocumentType, table: Table, col_widths: tuple[float, ...] | None = None) -> None:
    _write_table(document, table.columns, table.rows, caption=table.caption, col_widths=col_widths)


def _write_identity_table(document: DocumentType, pairs: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in pairs:
        row = table.add_row()
        _fill_cell(row.cells[0], label, bold=True, size=10.5)
        _shade(row.cells[0], HEADER_FILL)
        _fill_cell(row.cells[1], value, size=10.5)
    _apply_widths(table, (0.28, 0.72), 2)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


# --- Entry point ------------------------------------------------------------------------------

def generate_modul_docx(content: ModulContent) -> bytes:
    """Renders a schema-version-2 Modul Ajar as a .docx file."""
    document = Document()
    _style_fonts(document)
    _setup_page(document)

    identity = content.identity
    fase = f" (Fase {identity.fase})" if identity.fase else ""
    total_minutes = identity.meetings * identity.minutes_per_meeting

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.add_run("MODUL AJAR KURIKULUM MERDEKA").bold = True

    title = document.add_paragraph()
    title_run = title.add_run(_clean(content.title.strip() or identity.topic))
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = INK

    _write_identity_table(document, [
        ("Mata pelajaran", identity.subject),
        ("Kelas / Fase", f"Kelas {identity.grade_level}{fase}"),
        ("Topik", identity.topic),
        ("Alokasi waktu", f"{identity.meetings} pertemuan x {identity.minutes_per_meeting} menit ({total_minutes} menit)"),
        ("Model pembelajaran", content.learning_model.name),
    ])

    # A. Informasi umum
    _write_part(document, "A. INFORMASI UMUM")
    if content.initial_competencies:
        _write_heading(document, "Kompetensi awal")
        _write_list(document, content.initial_competencies)
    if content.pancasila_profile:
        _write_heading(document, "Profil Pelajar Pancasila")
        _write_table(
            document,
            ["Dimensi", "Wujud dalam kegiatan"],
            [[item.dimension, item.how] for item in content.pancasila_profile],
            col_widths=(0.3, 0.7),
        )
    if content.facilities:
        _write_heading(document, "Sarana dan prasarana")
        _write_list(document, content.facilities)
    if content.target_students.strip():
        _write_heading(document, "Target peserta didik")
        _write_paragraph(document, content.target_students)
    _write_heading(document, "Model pembelajaran")
    _write_paragraph(document, f"{content.learning_model.name}. {content.learning_model.rationale}")

    # B. Komponen inti
    _write_part(document, "B. KOMPONEN INTI")
    _write_heading(document, "Tujuan pembelajaran")
    _write_list(
        document,
        [objective.text for objective in content.learning_objectives],
        numbered=True,
        suffixes=[" ".join(f"[{source}]" for source in objective.source_ids) for objective in content.learning_objectives],
    )
    if content.meaningful_understanding:
        _write_heading(document, "Pemahaman bermakna")
        _write_list(document, content.meaningful_understanding)
    if content.trigger_questions:
        _write_heading(document, "Pertanyaan pemantik")
        _write_list(document, content.trigger_questions)
    if content.preparation:
        _write_heading(document, "Persiapan pembelajaran")
        _write_list(document, content.preparation)

    _write_heading(document, "Kegiatan pembelajaran")
    for meeting in content.meetings:
        _write_label(document, f"Pertemuan {meeting.number}: {meeting.title}")
        widths = ACTIVITY_COLUMN_WIDTHS if len(meeting.activities.columns) == len(ACTIVITY_COLUMN_WIDTHS) else None
        _write_model_table(document, meeting.activities, widths)

    _write_heading(document, "Asesmen")
    for label, item in (
        ("Asesmen diagnostik", content.assessment.diagnostic),
        ("Asesmen formatif", content.assessment.formative),
        ("Asesmen sumatif", content.assessment.summative),
    ):
        _write_label(document, label)
        _write_paragraph(document, f"Teknik: {item.technique}. Instrumen: {item.instrument}. {item.description}")
    _write_model_table(document, content.assessment.overview)
    _write_model_table(document, content.assessment.rubric)

    if content.enrichment:
        _write_heading(document, "Pengayaan")
        _write_list(document, content.enrichment)
    if content.remedial:
        _write_heading(document, "Remedial")
        _write_list(document, content.remedial)
    if content.reflection.teacher or content.reflection.students:
        _write_heading(document, "Refleksi")
        if content.reflection.teacher:
            _write_label(document, "Untuk guru")
            _write_list(document, content.reflection.teacher)
        if content.reflection.students:
            _write_label(document, "Untuk peserta didik")
            _write_list(document, content.reflection.students)

    # C. Lampiran
    _write_part(document, "C. LAMPIRAN")
    worksheet = content.attachments.worksheet
    _write_heading(document, f"LKPD: {worksheet.title}")
    _write_paragraph(document, worksheet.instructions)
    _write_list(document, worksheet.tasks, numbered=True)
    if content.attachments.reading_material.strip():
        _write_heading(document, "Bahan bacaan peserta didik")
        for paragraph in content.attachments.reading_material.split("\n"):
            _write_paragraph(document, paragraph)
    if content.attachments.glossary:
        _write_heading(document, "Glosarium")
        _write_table(
            document,
            ["Istilah", "Pengertian"],
            [[entry.term, entry.definition] for entry in content.attachments.glossary],
            col_widths=(0.3, 0.7),
        )

    if content.references:
        _write_heading(document, "Sumber rujukan")
        _write_list(document, [reference_line(reference) for reference in content.references])

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
